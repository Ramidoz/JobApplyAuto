"""
generate_docx.py — Produces formatted .docx files from a JSON spec.

Usage:
  python generate_docx.py --spec <path_to_spec.json> --out <output.docx> [--type resume|coverletter]

Spec format for resume:
{
  "candidate": {
    "name": "Rohit Ananthan",
    "tagline": "Data Scientist  |  Machine Learning  |  Product Analytics  |  MLOps",
    "contact": "+1 (202) 989-9596  |  rohitananthan123@gmail.com  |  linkedin.com/in/rohit-ananthan  |  github.com/Ramidoz"
  },
  "profile": "Tailored profile paragraph...",
  "skills": [
    {"label": "Languages", "value": "Python (pandas, NumPy, scikit-learn, PySpark), SQL, R"},
    ...
  ],
  "experience": [
    {
      "title": "Data Scientist Consultant",
      "company": "Invision Global Tech Inc.",
      "dates": "Feb 2026 – Present",
      "location": "United States, Remote",
      "bullets": ["Bullet 1", "Bullet 2"]
    }
  ],
  "education": [
    {
      "degree": "M.S. Information Systems",
      "school": "University of Maryland, College Park",
      "dates": "Dec 2024",
      "notes": "Relevant: Machine Learning, Data Engineering, Statistical Analysis, Database Systems"
    }
  ],
  "certifications": ["AWS Certified AI Practitioner (AIF-C01) — Issued Feb 2025"]
}

Spec format for cover letter:
{
  "candidate": {
    "name": "Rohit Ananthan",
    "contact": "rohitananthan123@gmail.com  |  linkedin.com/in/rohit-ananthan  |  rohitananthan.info"
  },
  "date": "March 11, 2026",
  "company": "Acme Corp",
  "role": "Senior Data Scientist",
  "body": "Full cover letter text..."
}
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("ERROR: python-docx not installed. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)


def set_font(run, name="Calibri", size=10, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_horizontal_rule(doc):
    """Add a thin horizontal line (border bottom on an empty paragraph)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "2E74B5")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def set_margins(doc, top=0.4, bottom=0.4, left=0.6, right=0.6):
    section = doc.sections[0]
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)


def build_resume(doc, spec):
    set_margins(doc)
    candidate = spec["candidate"]

    # Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(candidate["name"])
    set_font(run, size=18, bold=True, color=(44, 116, 181))

    # Tagline
    if candidate.get("tagline"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(candidate["tagline"])
        set_font(run, size=9, color=(89, 89, 89))

    # Contact
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(candidate["contact"])
    set_font(run, size=9, color=(89, 89, 89))

    # Profile
    add_section_header(doc, "PROFILE")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(spec["profile"])
    set_font(run, size=9.5)

    # Core Skills
    add_section_header(doc, "CORE SKILLS")
    for skill in spec.get("skills", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run_label = p.add_run(skill["label"] + ": ")
        set_font(run_label, size=9, bold=True)
        run_val = p.add_run(skill["value"])
        set_font(run_val, size=9)

    # Experience
    add_section_header(doc, "EXPERIENCE")
    for job in spec.get("experience", []):
        # Title + dates on same line (tab separated)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run_title = p.add_run(job["title"])
        set_font(run_title, size=9.5, bold=True)
        run_dates = p.add_run("\t" + job["dates"])
        set_font(run_dates, size=9, color=(89, 89, 89))
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_ALIGN_PARAGRAPH.RIGHT)

        # Company + location
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run_co = p.add_run(job["company"] + "  |  " + job["location"])
        set_font(run_co, size=9, color=(89, 89, 89))

        # Bullets
        for bullet in job.get("bullets", []):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(bullet)
            set_font(run, size=9)

    # Education & Certifications
    add_section_header(doc, "EDUCATION & CERTIFICATIONS")
    for edu in spec.get("education", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run_school = p.add_run(edu["school"])
        set_font(run_school, size=9.5, bold=True)
        run_dates = p.add_run("\t" + edu["dates"])
        set_font(run_dates, size=9, color=(89, 89, 89))
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.2), WD_ALIGN_PARAGRAPH.RIGHT)

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(edu["degree"])
        set_font(run, size=9)
        if edu.get("notes"):
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(1)
            run2 = p2.add_run(edu["notes"])
            set_font(run2, size=8.5, color=(89, 89, 89))

    for cert in spec.get("certifications", []):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(cert)
        set_font(run, size=9)


def add_section_header(doc, text):
    add_horizontal_rule(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_font(run, size=10, bold=True, color=(44, 116, 181))


def build_coverletter(doc, spec):
    set_margins(doc, top=1.0, bottom=1.0, left=1.0, right=1.0)
    candidate = spec["candidate"]

    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(candidate["name"])
    set_font(run, size=14, bold=True, color=(44, 116, 181))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(candidate["contact"])
    set_font(run, size=9.5, color=(89, 89, 89))

    # Date + company
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(spec.get("date", ""))
    set_font(run, size=10)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(spec.get("company", "") + " — " + spec.get("role", ""))
    set_font(run, size=10, bold=True)

    # Body (split on double newlines for paragraphs)
    body = spec.get("body", "")
    paragraphs = [b.strip() for b in body.split("\n\n") if b.strip()]
    for i, para_text in enumerate(paragraphs):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(para_text)
        set_font(run, size=10.5)


def main():
    parser = argparse.ArgumentParser(description="Generate .docx from JSON spec")
    parser.add_argument("--spec", required=True, help="Path to JSON spec file")
    parser.add_argument("--out", required=True, help="Path for output .docx file")
    parser.add_argument("--type", choices=["resume", "coverletter"], default="resume",
                        help="Document type (default: resume)")
    args = parser.parse_args()

    spec_path = Path(args.spec)
    out_path = Path(args.out)

    if not spec_path.exists():
        print(f"ERROR: Spec file not found: {spec_path}", file=sys.stderr)
        sys.exit(1)

    out_path.parent.mkdir(parents=True, exist_ok=True)

    with open(spec_path, "r", encoding="utf-8") as f:
        spec = json.load(f)

    doc = Document()
    # Remove default styles that add unwanted spacing
    for style in doc.styles:
        if hasattr(style, "paragraph_format"):
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)

    if args.type == "resume":
        build_resume(doc, spec)
    else:
        build_coverletter(doc, spec)

    doc.save(str(out_path))
    print(f"OK: {out_path}")


if __name__ == "__main__":
    main()
